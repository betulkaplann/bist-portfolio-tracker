"""
tracker.py - core functions for the BIST portfolio tracker.

The notebook (bist_tracker.ipynb) and the script (run_report.py) both import
these functions, so every calculation lives in ONE place.

Workflow:
    1. load_holdings()        -> read your positions from a CSV file
    2. get_prices()           -> download daily closing prices (Yahoo Finance)
    3. portfolio_value()      -> daily value of the portfolio
    4. holdings_table()       -> current value, profit/loss and weight per stock
    5. performance_metrics()  -> return, volatility, Sharpe, drawdown, beta ...
    6. export_excel()         -> Excel report with live formulas and a chart
"""

from pathlib import Path

import numpy as np
import pandas as pd

# On a Mac, Python from python.org often cannot check website certificates
# ("CERTIFICATE_VERIFY_FAILED"). truststore makes Python use the Mac's own
# certificate store, which fixes this for all data sources.
try:
    import truststore
    truststore.inject_into_ssl()
except ImportError:
    pass

# Number of trading days in a year (used to annualise returns and volatility)
TRADING_DAYS = 252

# Yahoo Finance codes for the benchmark and the exchange rate
BENCHMARK = "XU100.IS"   # BIST 100 index
FX_TICKER = "TRY=X"      # how many TRY one USD buys (USD/TRY)

# Folder of this file, so paths work wherever the project is saved
PROJECT_DIR = Path(__file__).resolve().parent
DATA_DIR = PROJECT_DIR / "data"


# ---------------------------------------------------------------------------
# 1. Holdings
# ---------------------------------------------------------------------------
def load_holdings(path=DATA_DIR / "holdings_sample.csv"):
    """Read the holdings file.

    Columns: ticker, sector, shares, amount_try, buy_price, buy_date
      - Tickers are written WITHOUT '.IS' (e.g. THYAO); we add it for Yahoo.
      - Give EITHER shares OR amount_try (TRY invested); the other is filled in.
      - sector is optional and used for the diversification analysis.
    """
    holdings = pd.read_csv(path, parse_dates=["buy_date"])
    holdings["ticker"] = holdings["ticker"].str.upper().str.strip()
    holdings["yahoo"] = holdings["ticker"] + ".IS"          # THYAO -> THYAO.IS
    for col in ["shares", "amount_try", "buy_price"]:       # optional columns
        if col not in holdings:
            holdings[col] = np.nan
    if "sector" not in holdings:
        holdings["sector"] = "Unknown"
    return holdings


def fill_buy_prices(holdings, prices):
    """If buy_price is empty in the CSV, use the closing price on buy_date.

    Handy for the sample portfolio. For your real portfolio, type the price
    you actually paid into the CSV instead.
    """
    holdings = holdings.copy()
    for i, row in holdings.iterrows():
        if pd.isna(row["buy_price"]):
            # first trading day on or after the buy date
            on_or_after = prices.loc[prices.index >= row["buy_date"], row["yahoo"]]
            holdings.loc[i, "buy_price"] = round(float(on_or_after.dropna().iloc[0]), 2)
    # If only an amount in TRY was given, buy as many whole shares as it allows
    no_shares = holdings["shares"].isna()
    holdings.loc[no_shares, "shares"] = np.floor(
        holdings.loc[no_shares, "amount_try"] / holdings.loc[no_shares, "buy_price"])
    holdings["cost"] = holdings["shares"] * holdings["buy_price"]
    return holdings


# ---------------------------------------------------------------------------
# 2. Prices
# ---------------------------------------------------------------------------
SOURCES = ["yahoo", "borsapy", "isyatirim"]   # tried in this order


def _clean_index(frame):
    """Make all dates plain days (no time zone, no hours) so sources match."""
    frame.index = pd.to_datetime(frame.index)
    if frame.index.tz is not None:
        frame.index = frame.index.tz_localize(None)
    frame.index = frame.index.normalize()
    frame.index.name = "Date"
    return frame[~frame.index.duplicated(keep="last")]


def _from_yahoo(tickers, start, end):
    """Yahoo Finance via the yfinance package."""
    import yfinance as yf

    raw = yf.download(tickers, start=start, end=end, auto_adjust=True,
                      progress=False)
    prices = raw["Close"]
    if isinstance(prices, pd.Series):                  # only one ticker
        prices = prices.to_frame(tickers[0])
    return prices


def _from_borsapy(tickers, start, end):
    """borsapy package: stocks/index from TradingView, USD/TRY from doviz.com."""
    import borsapy as bp

    columns = {}
    for t in tickers:
        try:
            if t == BENCHMARK:                                  # BIST 100 index
                hist = bp.Index("XU100").history(start=start, end=end)
                series = hist["Close"]
            elif t == FX_TICKER:                                # USD/TRY
                hist = bp.FX("USD").history(start=start, end=end)
                series = hist["Close"]
            else:                                               # a stock
                hist = bp.Ticker(t.replace(".IS", "")).history(
                    start=start, end=end, auto_adjust=True)
                # "Adj Close" also corrects for dividends, like Yahoo
                series = hist["Adj Close"] if "Adj Close" in hist else hist["Close"]
            columns[t] = _clean_index(series.to_frame())[series.name]
        except Exception as error:
            print(f"  borsapy: no data for {t} ({error})")
    return pd.DataFrame(columns)


def _from_isyatirim(tickers, start, end):
    """isyatirimhisse package: stocks and BIST 100 from the Is Yatirim website."""
    from isyatirimhisse import fetch_index_data, fetch_stock_data

    fmt = lambda d: pd.Timestamp(d).strftime("%d-%m-%Y")   # this API wants dd-mm-yyyy
    end_str = fmt(end) if end else None
    columns = {}

    stocks = [t for t in tickers if t.endswith(".IS") and t != BENCHMARK]
    if stocks:
        raw = fetch_stock_data([t.replace(".IS", "") for t in stocks],
                               start_date=fmt(start), end_date=end_str)
        # closing price column (adjusted if available)
        close_col = next((c for c in ["HGDG_KAPANIS", "HG_KAPANIS"] if c in raw), None)
        if close_col is None:
            raise KeyError(f"no closing-price column; columns are {list(raw.columns)}")
        wide = raw.pivot_table(index="HGDG_TARIH", columns="HGDG_HS_KODU",
                               values=close_col)
        for t in stocks:
            if t.replace(".IS", "") in wide:
                columns[t] = wide[t.replace(".IS", "")]

    if BENCHMARK in tickers:
        idx = fetch_index_data("XU100", start_date=fmt(start), end_date=end_str)
        columns[BENCHMARK] = idx.set_index("DATE")["VALUE"]
    # USD/TRY is not available here; it is filled in by another source
    prices = pd.DataFrame(columns)
    return _clean_index(prices) if not prices.empty else prices


def get_prices(tickers, start, end=None, offline=False, sources=SOURCES):
    """Daily closing prices, one column per ticker.

    Tries each data source in turn (Yahoo -> borsapy -> Is Yatirim). If one
    source is missing some tickers, the next source fills the gaps.
    If nothing works (e.g. no internet), or offline=True, it reads the
    SIMULATED prices in data/sample_prices.csv so the code still runs.
    """
    fetchers = {"yahoo": _from_yahoo, "borsapy": _from_borsapy,
                "isyatirim": _from_isyatirim}
    prices = pd.DataFrame()
    used = []

    if not offline:
        for name in sources:
            missing = [t for t in tickers if t not in prices
                       or prices[t].dropna().empty]
            if not missing:
                break
            try:
                new = fetchers[name](missing, start, end)
                new = new.dropna(how="all", axis=1)
                if new.empty:
                    print(f"{name}: no data")
                    continue
                new = _clean_index(new)
                prices = new if prices.empty else prices.combine_first(new)
                used.append(f"{name} ({', '.join(new.columns)})")
                print(f"{name}: got {len(new.columns)} of {len(missing)} series")
            except Exception as error:          # package missing, no internet ...
                print(f"{name}: failed ({type(error).__name__}: {error})")

        missing = [t for t in tickers if t not in prices or prices[t].dropna().empty]
        if not missing:
            prices = prices[tickers].sort_index()
            prices.attrs["source"] = "; ".join(used)
            return prices
        print(f"Still missing {missing} - using simulated sample prices instead.")

    prices = pd.read_csv(DATA_DIR / "sample_prices.csv",
                         index_col="Date", parse_dates=True)
    prices = prices.loc[prices.index >= pd.Timestamp(start), tickers]
    prices.attrs["source"] = "SIMULATED sample data (not real market prices)"
    return prices


# ---------------------------------------------------------------------------
# 3. Portfolio value over time
# ---------------------------------------------------------------------------
def portfolio_value(holdings, prices):
    """Daily market value of the portfolio in TRY.

    For each day: value = sum over stocks of (shares x closing price).
    A stock only counts from its buy_date onwards.
    """
    value = pd.DataFrame(index=prices.index)
    for _, row in holdings.iterrows():
        shares_held = np.where(prices.index >= row["buy_date"], row["shares"], 0)
        value[row["ticker"]] = shares_held * prices[row["yahoo"]]
    # ffill: if one stock has no price on a day (holiday/halt), use the last one
    value = value.ffill()
    return value.sum(axis=1).rename("Portfolio")


def portfolio_index(holdings, prices):
    """Time-weighted performance index (starts at 100).

    Why not just use the portfolio value? When you BUY a new stock, the value
    jumps up because new money comes in - that is not investment return.
    The time-weighted return removes these cash flows:

        daily return = (value today - new money today) / value yesterday - 1

    This is the standard way fund managers report performance.
    """
    value = portfolio_value(holdings, prices)
    first_day = value[value > 0].index[0]

    # New money on each day = shares bought x price on the buy day
    flows = pd.Series(0.0, index=prices.index)
    for _, row in holdings.iterrows():
        buy_day = prices.index[prices.index >= row["buy_date"]][0]
        flows[buy_day] += row["shares"] * prices.loc[buy_day, row["yahoo"]]

    value, flows = value.loc[first_day:], flows.loc[first_day:]
    daily = (value - flows) / value.shift(1) - 1
    daily.iloc[0] = 0.0                               # first day: no return yet
    return (100 * (1 + daily).cumprod()).rename("Portfolio")


def holdings_table(holdings, prices):
    """One row per stock: last price, market value, profit/loss and weight."""
    last_prices = prices.ffill().iloc[-1]
    table = holdings[["ticker", "sector", "shares", "buy_price", "cost"]].copy()
    table["last_price"] = table["ticker"].map(lambda t: last_prices[t + ".IS"])
    table["market_value"] = table["shares"] * table["last_price"]
    table["pnl"] = table["market_value"] - table["cost"]
    table["pnl_pct"] = table["pnl"] / table["cost"]
    table["weight"] = table["market_value"] / table["market_value"].sum()
    return table.sort_values("weight", ascending=False).reset_index(drop=True)


# ---------------------------------------------------------------------------
# 4. Performance metrics
# ---------------------------------------------------------------------------
def max_drawdown(series):
    """Largest fall from a previous peak, e.g. -0.25 = -25%."""
    running_peak = series.cummax()
    drawdown = series / running_peak - 1
    return drawdown.min()


def performance_metrics(value, benchmark, risk_free_annual=0.0):
    """Key risk/return numbers for a value series vs a benchmark series.

    risk_free_annual: yearly risk-free rate as a decimal (e.g. 0.35 = 35%).
    For TRY returns use a Turkish rate (e.g. the TCMB policy rate or a
    1-year government bond yield) - check the current level yourself.
    """
    # Line the two series up on the same dates
    data = pd.concat([value, benchmark], axis=1, join="inner").dropna()
    data.columns = ["portfolio", "benchmark"]

    daily = data.pct_change().dropna()                 # daily % changes
    years = len(daily) / TRADING_DAYS

    total_return = data["portfolio"].iloc[-1] / data["portfolio"].iloc[0] - 1
    bench_return = data["benchmark"].iloc[-1] / data["benchmark"].iloc[0] - 1
    cagr = (1 + total_return) ** (1 / years) - 1       # average yearly growth
    volatility = daily["portfolio"].std() * np.sqrt(TRADING_DAYS)

    # Sharpe ratio: extra return per unit of risk
    rf_daily = (1 + risk_free_annual) ** (1 / TRADING_DAYS) - 1
    excess = daily["portfolio"] - rf_daily
    sharpe = excess.mean() / excess.std() * np.sqrt(TRADING_DAYS)

    # Beta: how strongly the portfolio moves with the BIST 100
    beta = daily.cov().loc["portfolio", "benchmark"] / daily["benchmark"].var()

    # Tracking error and information ratio: how far we deviate from the index
    active = daily["portfolio"] - daily["benchmark"]
    tracking_error = active.std() * np.sqrt(TRADING_DAYS)
    info_ratio = active.mean() * TRADING_DAYS / tracking_error

    return pd.Series({
        "Start date": data.index[0].date(),
        "End date": data.index[-1].date(),
        "Total return": total_return,
        "Benchmark return": bench_return,
        "Excess return": total_return - bench_return,
        "CAGR": cagr,
        "Annual volatility": volatility,
        "Sharpe ratio": sharpe,
        "Max drawdown": max_drawdown(data["portfolio"]),
        "Benchmark max drawdown": max_drawdown(data["benchmark"]),
        "Beta vs BIST 100": beta,
        "Tracking error": tracking_error,
        "Information ratio": info_ratio,
    })


def to_usd(series_try, usdtry):
    """Convert a TRY series to USD (divide by the USD/TRY rate each day)."""
    fx = usdtry.reindex(series_try.index).ffill()
    return (series_try / fx).rename(series_try.name)


# ---------------------------------------------------------------------------
# 5. Charts
# ---------------------------------------------------------------------------
def make_charts(value_try, bench_try, value_usd, bench_usd, table, out_dir,
                source=""):
    """Save four PNG charts to out_dir and return their file paths."""
    import matplotlib
    matplotlib.use("Agg")                     # draw to files, no window needed
    import matplotlib.pyplot as plt

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    note = f"Source: {source}" if source else ""
    paths = []

    def indexed(a, b):
        both = pd.concat([a, b], axis=1, join="inner").dropna()
        return both / both.iloc[0] * 100

    # (a) Growth of 100 in TRY and in USD, side by side
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.5), sharex=True)
    for ax, (a, b, cur) in zip(axes, [(value_try, bench_try, "TRY"),
                                      (value_usd, bench_usd, "USD")]):
        idx = indexed(a, b)
        ax.plot(idx.index, idx.iloc[:, 0], label="Portfolio", linewidth=2)
        ax.plot(idx.index, idx.iloc[:, 1], label="BIST 100", linewidth=1.5)
        ax.axhline(100, color="grey", linewidth=0.8, linestyle="--")
        ax.set_title(f"Growth of 100 ({cur})")
        ax.legend()
        ax.grid(alpha=0.3)
    fig.text(0.01, 0.01, note, fontsize=8, color="grey")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    paths.append(out_dir / "growth_try_usd.png")
    fig.savefig(paths[-1], dpi=120)
    plt.close(fig)

    # (b) Drawdown: how far below its previous peak the portfolio is
    fig, ax = plt.subplots(figsize=(12, 3.5))
    for series, label in [(value_try, "Portfolio"), (bench_try, "BIST 100")]:
        dd = series / series.cummax() - 1
        ax.fill_between(dd.index, dd * 100, 0, alpha=0.35, label=label)
    ax.set_title("Drawdown from previous peak (%, TRY)")
    ax.legend()
    ax.grid(alpha=0.3)
    fig.text(0.01, 0.01, note, fontsize=8, color="grey")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    paths.append(out_dir / "drawdown.png")
    fig.savefig(paths[-1], dpi=120)
    plt.close(fig)

    # (c) Allocation and P/L per stock
    fig, axes = plt.subplots(1, 2, figsize=(12, 4))
    t = table.sort_values("weight")
    axes[0].barh(t["ticker"], t["weight"] * 100)
    axes[0].set_title("Portfolio weight (%)")
    colors = ["tab:green" if x >= 0 else "tab:red" for x in t["pnl_pct"]]
    axes[1].barh(t["ticker"], t["pnl_pct"] * 100, color=colors)
    axes[1].set_title("Profit / loss since purchase (%, TRY)")
    for ax in axes:
        ax.grid(alpha=0.3, axis="x")
    fig.text(0.01, 0.01, note, fontsize=8, color="grey")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    paths.append(out_dir / "allocation_pnl.png")
    fig.savefig(paths[-1], dpi=120)
    plt.close(fig)

    # (d) Rolling 60-day volatility (annualised): is risk rising or falling?
    fig, ax = plt.subplots(figsize=(12, 3.5))
    for series, label in [(value_try, "Portfolio"), (bench_try, "BIST 100")]:
        vol = series.pct_change().rolling(60).std() * np.sqrt(TRADING_DAYS) * 100
        ax.plot(vol.index, vol, label=label)
    ax.set_title("Rolling 60-day volatility (% per year, TRY)")
    ax.legend()
    ax.grid(alpha=0.3)
    fig.text(0.01, 0.01, note, fontsize=8, color="grey")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    paths.append(out_dir / "rolling_volatility.png")
    fig.savefig(paths[-1], dpi=120)
    plt.close(fig)

    return paths


# ---------------------------------------------------------------------------
# 6. Excel report
# ---------------------------------------------------------------------------
def export_excel(table, metrics_try, metrics_usd, value, benchmark, path,
                 source=""):
    """Write an Excel report.

    The Holdings sheet uses real Excel FORMULAS for market value, P/L and
    weight, so you can type a new price into the workbook and it recalculates.
    """
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="1F3864")
    blue_input = Font(color="0000FF")          # blue = number you can change

    def style_header(ws, row, n_cols):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

    # ---- Sheet 1: Holdings (with formulas) -------------------------------
    ws = wb.active
    ws.title = "Holdings"
    headers = ["Ticker", "Shares", "Buy price (TRY)", "Last price (TRY)",
               "Cost (TRY)", "Market value (TRY)", "P/L (TRY)", "P/L %", "Weight"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    first, last = 2, len(table) + 1
    for i, row in enumerate(table.itertuples(), start=first):
        ws.cell(row=i, column=1, value=row.ticker)
        ws.cell(row=i, column=2, value=float(row.shares)).font = blue_input
        ws.cell(row=i, column=3, value=float(row.buy_price)).font = blue_input
        ws.cell(row=i, column=4, value=round(float(row.last_price), 2)).font = blue_input
        ws.cell(row=i, column=5, value=f"=B{i}*C{i}")
        ws.cell(row=i, column=6, value=f"=B{i}*D{i}")
        ws.cell(row=i, column=7, value=f"=F{i}-E{i}")
        ws.cell(row=i, column=8, value=f"=G{i}/E{i}")
        ws.cell(row=i, column=9, value=f"=F{i}/$F${last + 1}")
    total = last + 1
    ws.cell(row=total, column=1, value="TOTAL").font = Font(bold=True)
    for col in "EFG":
        ws[f"{col}{total}"] = f"=SUM({col}{first}:{col}{last})"
        ws[f"{col}{total}"].font = Font(bold=True)
    ws[f"H{total}"] = f"=G{total}/E{total}"
    ws[f"I{total}"] = f"=SUM(I{first}:I{last})"
    for r in range(first, total + 1):
        for col in "BCDEFG":
            ws[f"{col}{r}"].number_format = "#,##0.00"
        for col in "HI":
            ws[f"{col}{r}"].number_format = "0.0%"
    ws.cell(row=total + 2, column=1,
            value="Blue numbers are inputs - change them and the formulas update.")
    ws.cell(row=total + 3, column=1, value=f"Price source: {source}")

    # ---- Sheet 2: Metrics -------------------------------------------------
    wm = wb.create_sheet("Metrics")
    wm.append(["Metric", "In TRY", "In USD"])
    style_header(wm, 1, 3)
    pct_rows = {"Total return", "Benchmark return", "Excess return", "CAGR",
                "Annual volatility", "Max drawdown", "Benchmark max drawdown",
                "Tracking error"}
    for name in metrics_try.index:
        v_try, v_usd = metrics_try[name], metrics_usd[name]
        if isinstance(v_try, (float, np.floating)):
            v_try, v_usd = float(v_try), float(v_usd)
        wm.append([name, v_try, v_usd])
        if "date" in name.lower():
            fmt = "yyyy-mm-dd"
        elif name in pct_rows:
            fmt = "0.0%"
        else:
            fmt = "0.00"
        wm.cell(row=wm.max_row, column=2).number_format = fmt
        wm.cell(row=wm.max_row, column=3).number_format = fmt

    # ---- Sheet 3: Daily values + line chart -------------------------------
    wd = wb.create_sheet("Daily values")
    wd.append(["Date", "Portfolio (indexed 100)", "BIST 100 (indexed 100)"])
    style_header(wd, 1, 3)
    both = pd.concat([value, benchmark], axis=1, join="inner").dropna()
    indexed = both / both.iloc[0] * 100
    for date, (p, b) in indexed.iterrows():
        wd.append([date.date(), round(float(p), 2), round(float(b), 2)])
        wd.cell(row=wd.max_row, column=1).number_format = "yyyy-mm-dd"

    chart = LineChart()
    chart.title = "Portfolio vs BIST 100 (start = 100)"
    chart.y_axis.title = "Index"
    chart.height, chart.width = 9, 20
    data_ref = Reference(wd, min_col=2, max_col=3, min_row=1, max_row=wd.max_row)
    dates_ref = Reference(wd, min_col=1, min_row=2, max_row=wd.max_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(dates_ref)
    wm.add_chart(chart, "E2")

    # Column widths so everything is readable
    for sheet in (ws, wm, wd):
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 20

    wb.save(path)
    return path
